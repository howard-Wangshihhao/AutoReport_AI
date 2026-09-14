#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""依 Excel 的「判定結果 / 目前情況」產生廠商報告文字並回填 Report 欄位。

特色：
- Windows / Linux / macOS 可用。
- AI provider 支援 OpenAI、Gemini，以及直接 HTTP 連線的 Ollama。
- 只依賴 Python 套件，不需要安裝 Microsoft Excel 或 LibreOffice。
- 若來源檔含 x14 的清單型資料驗證，會先轉成 openpyxl 可維護的標準 DataValidation。

使用方式：
    python generate_report_ai.py check_iot_ai.xlsx

預設輸出：
    <input>_report.xlsx
"""

from __future__ import annotations

import argparse
from copy import copy
from dataclasses import dataclass
import json
import math
import os
import re
import sys
import unicodedata
import warnings
import zipfile
from pathlib import Path, PurePosixPath
from typing import Callable
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName


SYSTEM_PROMPT = """你是一位資安檢測報告與 CVSS v3.1 Base Metrics 輔助判讀人員。
你的輸入包含檢測工程師已決定的「判定結果」、測項名稱、「目前情況」與可能的修補建議。

必須遵守：
1. 判定結果是檢測工程師的既定結論，不得自行升級或降低。
2. report_en 只輸出英文正式報告；report_zh 只輸出繁體中文正式報告，兩者語意必須一致。
3. 報告敘述必須遵循指定的判定結果骨架與測項類別模板，不得新增未提供的測試事實。
4. 對 low / medium / high / critical risk，報告中的降低風險因素只能來自「目前情況」已明確提供的防護措施、限制條件或使用情境。
5. CVSS 欄位只評估 CVSS v3.1 Base Metrics：AV、AC、PR、UI、S、C、I、A。不得輸出 Temporal 或 Environmental metrics。
6. CVSS Base Metrics 描述漏洞本身的固有嚴重度，不應因特定部署環境的額外防火牆、ACL、白名單或其他環境緩解措施而降低 Base Metrics。
7. 對 low / medium / high / critical risk，請依「目前情況」與明確引用的參考測項內容，以 best-effort 方式完成八個 CVSS Base Metrics。不得任意編造測試事實，但不要僅因部分資訊未逐字寫出就直接回傳 null；應依漏洞型態、攻擊路徑與已提供條件做最合理判定。
8. not applicable 與 no risk 不需要產生標準 Base Metrics，cvss 輸出 null，由 Python 套用固定值。
9. 英文風險名稱一律小寫：not applicable、no risk、low risk、medium risk、high risk、critical risk。
10. 僅輸出一個 JSON object，不要 Markdown、code fence、前言或解釋。

JSON 格式：
{
  "report_en": "English report paragraph",
  "report_zh": "繁體中文報告段落",
  "cvss": {
    "AV": "N|A|L|P",
    "AC": "L|H",
    "PR": "N|L|H",
    "UI": "N|R",
    "S": "U|C",
    "C": "N|L|H",
    "I": "N|L|H",
    "A": "N|L|H"
  }
}

對 low / medium / high / critical risk，原則上必須輸出完整 cvss object；只有輸入內容完全不足以辨識漏洞或攻擊情境時才可設為 null。
"""

CVSS_RETRY_SYSTEM_PROMPT = """你是一位 CVSS v3.1 Base Metrics 輔助判讀人員。

你只需要針對已被人工判定為 low / medium / high / critical risk 的測項，依提供的測試事實重新判定完整的 CVSS v3.1 Base Metrics。

規則：
1. 只輸出 AV、AC、PR、UI、S、C、I、A 八個 Base Metrics。
2. 不得為了配合人工風險等級而刻意調整 metrics。
3. 不得把未提供的產品功能、攻擊結果、防護措施或漏洞細節寫成事實。
4. 若資訊未逐字寫出，可依漏洞型態、攻擊路徑、存取位置、是否需要使用者互動與實際影響做 best-effort 判定。
5. CVSS Base Metrics 描述漏洞固有特性，不應因額外防火牆、ACL、白名單等部署環境緩解措施而降低。
6. 若漏洞成功利用前，必須由攻擊者以外的合法使用者執行必要動作（例如按實體按鈕啟用配對），UI 應優先考慮 R；若攻擊者可自行觸發必要流程，則 UI=N。不要用 AC:H 取代 UI:R。
7. 僅在輸入內容完全無法辨識漏洞或攻擊情境時，才允許 cvss 為 null。
8. 不要計算 score；Python 會依 FIRST CVSS v3.1 公式計算。
9. 僅輸出 JSON object，不要 Markdown、code fence 或解釋。

JSON 格式：
{
  "cvss": {
    "AV": "N|A|L|P",
    "AC": "L|H",
    "PR": "N|L|H",
    "UI": "N|R",
    "S": "U|C",
    "C": "N|L|H",
    "I": "N|L|H",
    "A": "N|L|H"
  }
}
"""




CVSS_METRICS_SCHEMA = {
    "type": "object",
    "properties": {
        "AV": {"type": "string", "enum": ["N", "A", "L", "P"]},
        "AC": {"type": "string", "enum": ["L", "H"]},
        "PR": {"type": "string", "enum": ["N", "L", "H"]},
        "UI": {"type": "string", "enum": ["N", "R"]},
        "S": {"type": "string", "enum": ["U", "C"]},
        "C": {"type": "string", "enum": ["N", "L", "H"]},
        "I": {"type": "string", "enum": ["N", "L", "H"]},
        "A": {"type": "string", "enum": ["N", "L", "H"]},
    },
    "required": ["AV", "AC", "PR", "UI", "S", "C", "I", "A"],
    "additionalProperties": False,
}

REPORT_RESPONSE_SCHEMA = {
    "type": "object",
    "properties": {
        "report_en": {"type": "string"},
        "report_zh": {"type": "string"},
        "cvss": {
            "anyOf": [
                CVSS_METRICS_SCHEMA,
                {"type": "null"},
            ]
        },
    },
    "required": ["report_en", "report_zh", "cvss"],
    "additionalProperties": False,
}

CVSS_RETRY_RESPONSE_SCHEMA = {
    "type": "object",
    "properties": {
        "cvss": {
            "anyOf": [
                CVSS_METRICS_SCHEMA,
                {"type": "null"},
            ]
        }
    },
    "required": ["cvss"],
    "additionalProperties": False,
}

RECOMMENDATION_RESPONSE_SCHEMA = {
    "type": "object",
    "properties": {
        "recommendation_en": {"type": "string"},
        "recommendation_zh": {"type": "string"},
    },
    "required": ["recommendation_en", "recommendation_zh"],
    "additionalProperties": False,
}

REQUIRED_HEADERS = ["編號", "測項", "判定結果", "目前情況", "Report", "Report_ch", "cvss", "score"]
SKIP_RESULTS = {"tbd", "testing"}

X14_NS = "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"
XM_NS = "http://schemas.microsoft.com/office/excel/2006/main"
MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


def clean_text(value) -> str:
    if value is None:
        return ""
    return str(value).replace("\xa0", " ").strip()


def call_gemini_generate_content(generate_content_func: Callable[..., object], **kwargs):
    """呼叫 Gemini generate_content，隱藏 SDK 目前會誤顯示的 AFC 建議警告。"""
    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message=r"Direct use of automatic function calling \(AFC\) in Models\.generate_content.*",
        )
        return generate_content_func(**kwargs)


def _display_width_units(text: str) -> float:
    """估算文字在 Excel 儲存格中的顯示寬度；中日韓全形字約視為 2 個 ASCII 字元。"""
    units = 0.0
    for ch in clean_text(text):
        if ch == "\t":
            units += 4.0
        elif unicodedata.east_asian_width(ch) in {"W", "F"}:
            units += 2.0
        else:
            units += 1.0
    return units


def estimate_wrapped_lines(text: str, column_width: float | None) -> int:
    """依 Excel 欄寬與文字內容粗估自動換行後所需行數。"""
    value = clean_text(text)
    if not value:
        return 1

    # Excel 欄寬大致以字元為單位；保守留一些內距，避免估得太少。
    width = max(float(column_width or 10.0) - 1.5, 1.0)
    total_lines = 0
    for paragraph in value.splitlines() or [""]:
        units = _display_width_units(paragraph)
        total_lines += max(1, math.ceil(units / width))
    return max(total_lines, 1)


def auto_adjust_row_height(
    ws,
    row_idx: int,
    columns: list[int] | tuple[int, ...],
    *,
    min_height: float = 20.0,
    max_height: float = 405.0,
    line_height: float = 15.0,
    padding: float = 6.0,
    preserve_existing: bool = True,
) -> float:
    """依指定欄位的長文字估算列高，並啟用自動換行。

    Excel 的 AutoFit 不會由 openpyxl 自動計算，因此這裡以欄寬、字元寬度與
    明確換行數估算。若模板原本列高較大，預設保留原高度；只在內容需要時加高。
    """
    max_lines = 1

    for col_idx in columns:
        if not col_idx:
            continue
        cell = ws.cell(row_idx, col_idx)
        alignment = copy(cell.alignment)
        alignment.wrap_text = True
        alignment.vertical = "top"
        cell.alignment = alignment

        width = ws.column_dimensions[cell.column_letter].width
        max_lines = max(max_lines, estimate_wrapped_lines(cell.value, width))

    estimated = min(max_height, max(min_height, max_lines * line_height + padding))
    current = ws.row_dimensions[row_idx].height
    if preserve_existing and current is not None:
        estimated = min(max_height, max(estimated, float(current)))

    ws.row_dimensions[row_idx].height = estimated
    return estimated


def str_to_bool(value: str | None, default: bool = True) -> bool:
    if value is None:
        return default
    return value.strip().lower() in {"1", "true", "yes", "y", "on"}




@dataclass(frozen=True)
class OllamaConfig:
    host: str
    model: str
    timeout: float
    keep_alive: str
    structured_output: bool
    num_ctx: int
    think: bool

    @classmethod
    def from_env(cls) -> "OllamaConfig":
        host = os.getenv("OLLAMA_HOST", "").strip().rstrip("/")
        model = os.getenv("OLLAMA_MODEL", "").strip()
        if not host:
            raise RuntimeError(".env 尚未設定 OLLAMA_HOST，例如 http://192.168.50.241:11434")
        if not re.match(r"^https?://", host, re.IGNORECASE):
            raise RuntimeError("OLLAMA_HOST 必須包含 http:// 或 https://")
        if not model:
            raise RuntimeError(".env 尚未設定 OLLAMA_MODEL")

        raw_timeout = os.getenv("OLLAMA_TIMEOUT", "300").strip()
        try:
            timeout = float(raw_timeout)
        except ValueError as exc:
            raise RuntimeError(".env 的 OLLAMA_TIMEOUT 必須是數字") from exc
        if timeout <= 0:
            raise RuntimeError(".env 的 OLLAMA_TIMEOUT 必須大於 0")

        raw_num_ctx = os.getenv("OLLAMA_NUM_CTX", "8192").strip()
        try:
            num_ctx = int(raw_num_ctx)
        except ValueError as exc:
            raise RuntimeError(".env 的 OLLAMA_NUM_CTX 必須是整數") from exc
        if num_ctx <= 0:
            raise RuntimeError(".env 的 OLLAMA_NUM_CTX 必須大於 0")

        return cls(
            host=host,
            model=model,
            timeout=timeout,
            keep_alive=os.getenv("OLLAMA_KEEP_ALIVE", "10m").strip() or "10m",
            structured_output=str_to_bool(
                os.getenv("OLLAMA_STRUCTURED_OUTPUT", "true"), True
            ),
            num_ctx=num_ctx,
            think=str_to_bool(os.getenv("OLLAMA_THINK", "false"), False),
        )


class OllamaBackend:
    """直接透過 Ollama HTTP API 產生結構化回覆。"""

    def __init__(self, config: OllamaConfig | None = None, session=None):
        self.config = config or OllamaConfig.from_env()
        self.model = self.config.model
        self.timeout = self.config.timeout
        self.keep_alive = self.config.keep_alive
        self.structured_output = self.config.structured_output
        self.num_ctx = self.config.num_ctx
        self.think = self.config.think

        if session is None:
            try:
                import requests
            except ImportError as exc:
                raise RuntimeError(
                    "使用 AI_PROVIDER=ollama 需要 requests，請執行：pip install requests"
                ) from exc
            session = requests.Session()

        self.session = session
        # 內網 Ollama 通常不應經過系統 HTTP/HTTPS proxy。
        if hasattr(self.session, "trust_env"):
            self.session.trust_env = False

    def chat(
        self,
        *,
        system_prompt: str,
        user_prompt: str,
        schema: dict,
        temperature: float,
        label: str = "",
    ) -> str:
        payload = {
            "model": self.model,
            "messages": [
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_prompt},
            ],
            "stream": False,
            "think": self.think,
            "format": schema if self.structured_output else "json",
            "options": {
                "temperature": temperature,
                "num_ctx": self.num_ctx,
            },
            "keep_alive": self.keep_alive,
        }

        response = None
        try:
            response = self.session.post(
                f"{self.config.host}/api/chat",
                json=payload,
                timeout=self.timeout,
            )
            response.raise_for_status()
        except Exception as exc:
            detail = ""
            if response is not None:
                try:
                    error_payload = response.json()
                    if isinstance(error_payload, dict):
                        detail = clean_text(error_payload.get("error"))
                except Exception:
                    detail = clean_text(getattr(response, "text", ""))
            suffix = f"：{detail}" if detail else ""
            raise RuntimeError(
                f"Ollama API 呼叫失敗 ({self.config.host})；請確認主機可連線、Ollama 已啟動，"
                f"且模型 {self.model} 已存在{suffix}"
            ) from exc

        try:
            data = response.json()
        except Exception as exc:
            raise RuntimeError("Ollama API 回傳內容不是有效 JSON") from exc

        message = data.get("message") if isinstance(data, dict) else None
        content = clean_text(message.get("content")) if isinstance(message, dict) else ""
        if not content:
            raise RuntimeError("Ollama API 未回傳 message.content")

        total_ns = data.get("total_duration") if isinstance(data, dict) else None
        load_ns = data.get("load_duration") if isinstance(data, dict) else None
        prompt_count = data.get("prompt_eval_count") if isinstance(data, dict) else None
        eval_count = data.get("eval_count") if isinstance(data, dict) else None
        eval_ns = data.get("eval_duration") if isinstance(data, dict) else None

        perf_parts = []
        if isinstance(total_ns, (int, float)) and total_ns > 0:
            perf_parts.append(f"total: {total_ns / 1_000_000_000:.1f}s")
        if isinstance(load_ns, (int, float)) and load_ns > 0:
            perf_parts.append(f"load: {load_ns / 1_000_000_000:.1f}s")
        if isinstance(prompt_count, int):
            perf_parts.append(f"prompt: {prompt_count} tokens")
        if isinstance(eval_count, int):
            perf_parts.append(f"output: {eval_count} tokens")
        if (
            isinstance(eval_count, int)
            and isinstance(eval_ns, (int, float))
            and eval_ns > 0
        ):
            perf_parts.append(f"{eval_count / (eval_ns / 1_000_000_000):.1f} tok/s")
        if perf_parts:
            perf_label = clean_text(label) or "request"
            print(f"  [OLLAMA PERF] {perf_label} | " + " | ".join(perf_parts), flush=True)

        return content

    def close(self) -> None:
        if self.session is not None:
            try:
                self.session.close()
            except Exception:
                pass
            self.session = None


_OLLAMA_BACKEND: OllamaBackend | None = None


def get_ollama_backend() -> OllamaBackend:
    global _OLLAMA_BACKEND
    if _OLLAMA_BACKEND is None:
        _OLLAMA_BACKEND = OllamaBackend()
        print(
            f"[OLLAMA] {_OLLAMA_BACKEND.config.host} | model={_OLLAMA_BACKEND.model} "
            f"| ctx={_OLLAMA_BACKEND.num_ctx} | think={str(_OLLAMA_BACKEND.think).lower()}",
            flush=True,
        )
    return _OLLAMA_BACKEND


def close_ollama_backend() -> None:
    global _OLLAMA_BACKEND
    if _OLLAMA_BACKEND is not None:
        _OLLAMA_BACKEND.close()
        _OLLAMA_BACKEND = None


def should_process(result: str, current_status: str, overwrite_report: bool, existing_report: str) -> bool:
    result_text = clean_text(result)
    current_text = clean_text(current_status)
    existing_text = clean_text(existing_report)

    if not result_text or not current_text:
        return False
    if result_text.casefold() in SKIP_RESULTS:
        return False
    if not overwrite_report and existing_text:
        return False
    return True



RESULT_ALIASES = {
    "not_applicable": {"not applicable", "n/a", "na", "不適用"},
    "none": {"none", "no risk", "no finding", "無風險", "未發現風險"},
    "low": {"low", "low risk", "低風險"},
    "medium": {"medium", "medium risk", "中風險"},
    "high": {"high", "high risk", "高風險"},
    "critical": {"critical", "critical risk", "嚴重風險", "重大風險"},
}

RESULT_LABEL_ZH = {
    "not_applicable": "不適用",
    "none": "無風險",
    "low": "低風險",
    "medium": "中風險",
    "high": "高風險",
    "critical": "嚴重風險",
}

RESULT_LABEL_EN = {
    "not_applicable": "not applicable",
    "none": "no risk",
    "low": "low risk",
    "medium": "medium risk",
    "high": "high risk",
    "critical": "critical risk",
}


# 依測項名稱分類；順序很重要，較具體的規則放前面。
TEST_ITEM_CATEGORY_RULES = [
    (
        "software_analysis",
        (
            "source code", "source-code", "sbom", "vulnerability scanning",
            "static application analysis", "static analysis", "sca",
        ),
    ),
    (
        "authentication",
        (
            "authentication", "password", "credential", "login", "account",
        ),
    ),
    (
        "network",
        (
            "network security", "network service", "communication", "tls", "ssl",
            "bluetooth", "ble", "wireless", "remote service",
        ),
    ),
    (
        "data_protection",
        (
            "sensitive data", "data leakage", "privacy", "personal data",
            "data protection", "encryption at rest", "data storage",
        ),
    ),
    (
        "logging_monitoring",
        (
            "logging", "log ", "event notification", "security event", "audit",
            "monitoring", "notification",
        ),
    ),
    (
        "update_delivery",
        (
            "application update", "firmware update", "software update", "delivery protection",
            "application delivery", "update mechanism", "update",
        ),
    ),
    (
        "physical_general",
        (
            "physical protection", "physical security", "physical", "tamper",
            "local interface", "debug interface", "uart", "jtag",
        ),
    ),
    (
        "application_security",
        (
            "sandbox", "fuzzing", "input validation", "configuration management",
            "api input", "ui vulnerability", "device vulnerability", "api security",
        ),
    ),
]


TEST_ITEM_TEMPLATE_INSTRUCTIONS = {
    "software_analysis": """【測項類別：軟體／原始碼／元件弱點分析】
- 英文描述優先使用 source code, software component, SBOM, vulnerability scan, static analysis 等與實際測項相符的詞。
- 繁體中文描述優先使用「原始碼掃描」、「軟體元件」、「SBOM」、「弱點掃描」、「靜態分析」等與實際測項相符的詞。
- 若目前情況有工具名稱，可自然寫成 using / 透過該工具進行；未提供工具名稱則不得自行補工具。
- 若為不適用，優先說明缺少原始碼、SBOM、套件資訊或其他必要資料等「目前情況」實際提供的原因。""",

    "application_security": """【測項類別：應用程式與輸入處理安全】
- 依實際測項聚焦於應用程式行為、輸入處理、模糊測試、沙箱觀察、UI 弱點或設定管理。
- 英文可使用 application behavior, input validation, fuzzing, sandbox testing, configuration 等詞；繁中使用相對應專業詞。
- 不得因測項名稱自行宣稱已執行特定 payload、攻擊手法或工具。""",

    "authentication": """【測項類別：身分驗證與密碼管理】
- 聚焦帳號、登入、身分驗證、密碼強度、預設密碼、憑證或 API authentication 等實際內容。
- 英文敘述優先使用 authentication mechanism, login control, password policy, credential 等詞。
- 繁體中文優先使用「身分驗證機制」、「登入控制」、「密碼政策」、「帳號憑證」等詞。
- 風險降低因素如鎖定機制、實體操作、權限限制、多因素驗證等，只能在目前情況明確提供時使用。""",

    "network": """【測項類別：網路與通訊安全】
- 聚焦網路服務、通訊介面、遠端／鄰近攻擊面、連線保護、加密或非必要服務暴露。
- 英文優先使用 network service, communication channel, remote/adjacent access, encrypted communication 等與實際情況相符的詞。
- 繁體中文優先使用「網路服務」、「通訊通道」、「遠端／鄰近存取」、「加密傳輸」等詞。
- 不得自行推定通訊協定、連接埠、TLS 版本或加密演算法。""",

    "data_protection": """【測項類別：敏感資料與隱私保護】
- 聚焦敏感資料是否於儲存、傳輸、畫面、日誌或暫存位置暴露，以及既有保護措施。
- 英文優先使用 sensitive data, data exposure, data transmission/storage, protection mechanism 等詞。
- 繁體中文優先使用「敏感資料」、「資料暴露」、「資料傳輸／儲存」、「保護機制」等詞。
- 未明確提供資料類型、加密方式或保存位置時，不得自行補充。""",

    "logging_monitoring": """【測項類別：安全事件日誌與通知】
- 聚焦安全事件是否被記錄、日誌內容是否足以追蹤、異常是否通知，以及目前實際觀察結果。
- 英文優先使用 security event logging, audit record, event notification, monitoring 等詞。
- 繁體中文優先使用「安全事件日誌」、「稽核紀錄」、「事件通知」、「監控」等詞。
- 不得自行宣稱具有 SIEM、告警平台或保存期限。""",

    "update_delivery": """【測項類別：更新與軟體交付保護】
- 聚焦更新來源、更新完整性／真實性、版本控制、失敗處理或交付通道等目前情況實際提供內容。
- 英文優先使用 update mechanism, software delivery, integrity/authenticity verification, version 等詞。
- 繁體中文優先使用「更新機制」、「軟體交付」、「完整性／真實性驗證」、「版本」等詞。
- 不得自行假設有簽章、Secure Boot、回滾保護或 OTA。""",

    "physical_general": """【測項類別：實體與一般裝置安全】
- 聚焦實體接觸、外部介面、除錯介面、防拆措施或其他裝置層面的實際觀察。
- 英文優先使用 physical access, device interface, debug interface, physical protection 等與實際測項相符的詞。
- 繁體中文優先使用「實體存取」、「裝置介面」、「除錯介面」、「實體防護」等詞。
- 若測項實際內容與實體安全無關，仍以目前情況為最高優先，不要硬套詞彙。""",

    "general": """【測項類別：一般資安測試】
- 依測項名稱與目前情況，以中性、正式方式描述實際測試事實與判定依據。
- 不套用未被目前情況支持的專有名詞、工具、攻擊方式或防護措施。""",
}



CVSS_METRIC_ORDER = ("AV", "AC", "PR", "UI", "S", "C", "I", "A")
CVSS_ALLOWED_VALUES = {
    "AV": {"N", "A", "L", "P"},
    "AC": {"L", "H"},
    "PR": {"N", "L", "H"},
    "UI": {"N", "R"},
    "S": {"U", "C"},
    "C": {"N", "L", "H"},
    "I": {"N", "L", "H"},
    "A": {"N", "L", "H"},
}
CVSS_NO_RISK_VECTOR = "CVSS:3.1/AV:-/AC:-/PR:-/UI:-/S:-/C:N/I:N/A:N"


def parse_ai_payload(text: str) -> dict:
    """解析 AI 回傳的 JSON；允許模型意外包一層 code fence。"""
    raw = clean_text(text)
    if raw.startswith("```"):
        raw = raw.removeprefix("```json").removeprefix("```").strip()
        if raw.endswith("```"):
            raw = raw[:-3].strip()
    first = raw.find("{")
    last = raw.rfind("}")
    if first < 0 or last < first:
        raise ValueError("AI 回傳內容不是有效 JSON object")
    try:
        payload = json.loads(raw[first:last + 1])
    except json.JSONDecodeError as exc:
        raise ValueError(f"AI JSON 解析失敗：{exc}") from exc
    if not isinstance(payload, dict):
        raise ValueError("AI JSON 最外層必須是 object")
    report_en = clean_text(payload.get("report_en"))
    report_zh = clean_text(payload.get("report_zh"))
    if not report_en or not report_zh:
        raise ValueError("AI JSON 缺少 report_en 或 report_zh")
    payload["report_en"] = report_en
    payload["report_zh"] = report_zh
    return payload


def build_cvss_retry_prompt(
    *,
    item_no: str,
    test_item: str,
    result: str,
    current_status: str,
    remediation: str = "",
    reference_context: str = "",
) -> str:
    """第一輪 CVSS 缺失/格式錯誤時，使用更聚焦的 prompt 重試一次。"""
    return f"""請重新判定此測項的 CVSS v3.1 Base Metrics。

測項編號：{clean_text(item_no)}
測試項目：{clean_text(test_item)}
人工判定結果：{clean_text(result)}

目前情況：
{clean_text(current_status) or '(未提供)'}

修補建議：
{clean_text(remediation) or '(未提供)'}

參考測項內容：
{clean_text(reference_context) or '(無)'}

請以 best-effort 方式輸出完整 AV、AC、PR、UI、S、C、I、A。
不要輸出 score，也不要為了配合人工判定結果而調整 metrics。
僅輸出 JSON object。
"""


def parse_cvss_retry_payload(text: str) -> dict | None:
    """解析 CVSS 重試 API 的 JSON 回傳，只取 cvss object。"""
    raw = clean_text(text)
    if raw.startswith("```"):
        raw = raw.removeprefix("```json").removeprefix("```").strip()
        if raw.endswith("```"):
            raw = raw[:-3].strip()
    first = raw.find("{")
    last = raw.rfind("}")
    if first < 0 or last < first:
        raise ValueError("CVSS 重試回傳內容不是有效 JSON object")
    try:
        payload = json.loads(raw[first:last + 1])
    except json.JSONDecodeError as exc:
        raise ValueError(f"CVSS 重試 JSON 解析失敗：{exc}") from exc
    if not isinstance(payload, dict):
        raise ValueError("CVSS 重試 JSON 最外層必須是 object")
    metrics = payload.get("cvss")
    if metrics is None:
        return None
    if not isinstance(metrics, dict):
        raise ValueError("CVSS 重試欄位 cvss 必須是 object 或 null")
    return metrics


def normalize_cvss_metrics(metrics: dict | None) -> dict:
    if not isinstance(metrics, dict):
        raise ValueError("CVSS Base Metrics 資訊不足或格式錯誤")

    # 對 AI 常見輸出做容錯：key 不分大小寫，也接受完整英文名稱。
    source = {clean_text(k).upper(): v for k, v in metrics.items()}
    value_aliases = {
        "AV": {"NETWORK": "N", "ADJACENT": "A", "LOCAL": "L", "PHYSICAL": "P"},
        "AC": {"LOW": "L", "HIGH": "H"},
        "PR": {"NONE": "N", "LOW": "L", "HIGH": "H"},
        "UI": {"NONE": "N", "REQUIRED": "R"},
        "S": {"UNCHANGED": "U", "CHANGED": "C"},
        "C": {"NONE": "N", "LOW": "L", "HIGH": "H"},
        "I": {"NONE": "N", "LOW": "L", "HIGH": "H"},
        "A": {"NONE": "N", "LOW": "L", "HIGH": "H"},
    }

    normalized = {}
    for key in CVSS_METRIC_ORDER:
        raw = clean_text(source.get(key)).upper()
        value = value_aliases.get(key, {}).get(raw, raw)
        if value not in CVSS_ALLOWED_VALUES[key]:
            allowed = "/".join(sorted(CVSS_ALLOWED_VALUES[key]))
            raise ValueError(f"CVSS {key} 值無效：{value or '(空白)'}；允許值：{allowed}")
        normalized[key] = value
    return normalized


def cvss31_vector(metrics: dict) -> str:
    m = normalize_cvss_metrics(metrics)
    return "CVSS:3.1/" + "/".join(f"{k}:{m[k]}" for k in CVSS_METRIC_ORDER)


def cvss31_roundup(value: float) -> float:
    """依 FIRST CVSS v3.1 Appendix A 的 Roundup 實作。"""
    int_input = math.floor(value * 100000.0 + 0.5)
    if int_input % 10000 == 0:
        return int_input / 100000.0
    return (int_input // 10000 + 1) / 10.0


def cvss31_base_score(metrics: dict) -> float:
    m = normalize_cvss_metrics(metrics)
    av = {"N": 0.85, "A": 0.62, "L": 0.55, "P": 0.20}[m["AV"]]
    ac = {"L": 0.77, "H": 0.44}[m["AC"]]
    ui = {"N": 0.85, "R": 0.62}[m["UI"]]
    cia = {"H": 0.56, "L": 0.22, "N": 0.0}
    c, i, a = cia[m["C"]], cia[m["I"]], cia[m["A"]]

    if m["PR"] == "N":
        pr = 0.85
    elif m["PR"] == "L":
        pr = 0.68 if m["S"] == "C" else 0.62
    else:
        pr = 0.50 if m["S"] == "C" else 0.27

    iss = 1 - ((1 - c) * (1 - i) * (1 - a))
    if m["S"] == "U":
        impact = 6.42 * iss
    else:
        impact = 7.52 * (iss - 0.029) - 3.25 * ((iss - 0.02) ** 15)

    if impact <= 0:
        return 0.0

    exploitability = 8.22 * av * ac * pr * ui
    if m["S"] == "U":
        raw = min(impact + exploitability, 10.0)
    else:
        raw = min(1.08 * (impact + exploitability), 10.0)
    return cvss31_roundup(raw)


def cvss31_severity(score: float) -> str:
    if score == 0.0:
        return "none"
    if 0.1 <= score <= 3.9:
        return "low"
    if 4.0 <= score <= 6.9:
        return "medium"
    if 7.0 <= score <= 8.9:
        return "high"
    if 9.0 <= score <= 10.0:
        return "critical"
    raise ValueError(f"CVSS score 超出有效範圍：{score}")


def score_result_mismatch_message(result: str, score: float) -> str:
    """人工風險判定與 CVSS Base severity 不一致時只提示，不阻擋輸出。"""
    kind = normalize_result_kind(result)
    if kind not in {"low", "medium", "high", "critical"}:
        return ""
    severity = cvss31_severity(score)
    if severity == kind:
        return ""
    return (
        f"CVSS v3.1 計算結果為 {score:.1f} ({severity})，"
        f"與人工判定結果 {RESULT_LABEL_EN[kind]} 不一致；已保留人工判定並輸出 CVSS Base Score，請人工複核"
    )


def set_cvss_cells_mismatch_style(ws, row_idx: int, cvss_col: int, score_col: int, mismatch: bool) -> None:
    """CVSS severity 與人工判定不一致時，將 cvss / score 兩格字體標紅。"""
    if not mismatch:
        return
    for col_idx in (cvss_col, score_col):
        cell = ws.cell(row_idx, col_idx)
        font = copy(cell.font)
        font.color = "FFFF0000"
        cell.font = font


def copy_report_ch_style_to_report(ws, row_idx: int, report_col: int, report_ch_col: int) -> None:
    """將同列 Report_ch 的完整儲存格格式複製到 Report；不修改任何儲存格內容。"""
    report_cell = ws.cell(row_idx, report_col)
    report_ch_cell = ws.cell(row_idx, report_ch_col)
    report_cell.font = copy(report_ch_cell.font)
    report_cell.fill = copy(report_ch_cell.fill)
    report_cell.border = copy(report_ch_cell.border)
    report_cell.alignment = copy(report_ch_cell.alignment)
    report_cell.number_format = report_ch_cell.number_format
    report_cell.protection = copy(report_ch_cell.protection)


def format_output_cells(
    ws,
    row_idx: int,
    report_col: int,
    report_ch_col: int,
    cvss_col: int,
    score_col: int,
) -> None:
    """統一四個輸出欄位為 12pt、垂直置中、換行；保留既有字色（含 CVSS 紅字）。"""
    # Report 先沿用 Report_ch 的完整基礎格式，維持雙語欄位一致。
    copy_report_ch_style_to_report(ws, row_idx, report_col, report_ch_col)

    for col_idx in (report_col, report_ch_col, cvss_col, score_col):
        cell = ws.cell(row_idx, col_idx)

        # copy() 後再改 size，可保留原本字型名稱、粗斜體、以及紅字等顏色資訊。
        font = copy(cell.font)
        font.sz = 12
        cell.font = font

        alignment = copy(cell.alignment)
        alignment.vertical = "center"
        alignment.wrap_text = True
        cell.alignment = alignment


def cvss_for_result(result: str, metrics: dict | None) -> tuple[str, float | str]:
    kind = normalize_result_kind(result)
    if kind == "not_applicable":
        return "-", "-"
    if kind == "none":
        return CVSS_NO_RISK_VECTOR, 0.0
    if kind not in {"low", "medium", "high", "critical"}:
        raise ValueError(f"無法依判定結果產生 CVSS：{clean_text(result)}")

    normalized = normalize_cvss_metrics(metrics)
    vector = cvss31_vector(normalized)
    score = cvss31_base_score(normalized)
    return vector, score


def classify_test_item(test_item: str) -> str:
    """依測項名稱選擇描述模板；無法辨識時回傳 general。"""
    value = clean_text(test_item).casefold()
    for category, keywords in TEST_ITEM_CATEGORY_RULES:
        if any(keyword in value for keyword in keywords):
            return category
    return "general"


def test_item_template_instruction(test_item: str) -> str:
    category = classify_test_item(test_item)
    return TEST_ITEM_TEMPLATE_INSTRUCTIONS[category]


def normalize_result_kind(result: str) -> str:
    value = clean_text(result).casefold()
    for kind, aliases in RESULT_ALIASES.items():
        if value in aliases:
            return kind
    return "generic"


def is_chinese_language(language: str) -> bool:
    value = clean_text(language).casefold()
    return any(token in value for token in ("chinese", "zh", "中文", "繁體", "traditional chinese"))


def fixed_conclusion(result: str, language: str) -> str:
    kind = normalize_result_kind(result)
    if kind == "generic":
        if is_chinese_language(language):
            return f"因此，此測項評估為{clean_text(result)}。"
        return f"Therefore, this test item is assessed as {clean_text(result)}."

    if is_chinese_language(language):
        return f"因此，此測項評估為{RESULT_LABEL_ZH[kind]}。"
    return f"Therefore, this test item is assessed as {RESULT_LABEL_EN[kind]}."


def report_structure_instruction(result: str, language: str) -> str:
    kind = normalize_result_kind(result)
    conclusion = fixed_conclusion(result, language)

    if kind == "not_applicable":
        return f"""【固定敘述結構：不適用】
- 先依「目前情況」說明為何本測項無法或不需執行，例如廠商未提供必要資料、功能不存在或不在產品適用範圍。
- 接著明確說明本測項不在本次測試範圍內。
- 不要描述不存在的測試工具、測試結果或風險。
- 段落最後必須逐字使用：{conclusion}
- 參考句型：此測項因為【目前情況所述原因】，故不在本次測試範圍內，{conclusion}"""

    if kind == "none":
        return f"""【固定敘述結構：無風險】
- 先依「目前情況」描述實際執行的測試方式、工具、觀察內容或確認方式；只寫有提供的資訊。
- 再描述實際測試結果，例如未發現異常、未發現弱點或符合預期行為；不得自行補測試結果。
- 段落最後必須逐字使用：{conclusion}
- 參考句型：透過【目前情況中的測試方式】進行測試，【目前情況中的測試結果】，{conclusion}"""

    if kind in {"low", "medium", "high", "critical"}:
        return f"""【固定敘述結構：風險項目】
- 第一句先描述觀察到的風險：依「目前情況」說明發現的弱點、缺失、可能造成的影響或攻擊面，不得擴大解讀。
- 第二部分再描述降低風險或限制風險的原因，但只能引用「目前情況」已明確提供的防護措施、限制條件、使用情境或其他事實。
- 如果「目前情況」沒有任何降低風險的因素，不得自行創造，可省略降低風險句，或客觀寫明目前資訊未提供其他降低風險因素。
- 若有「修補建議」，可在固定結論前簡短整理改善方向。
- 段落最後必須逐字使用：{conclusion}"""

    return f"""【固定敘述結構：其他判定】
- 依「目前情況」客觀描述測試事實與判定依據。
- 不得自行新增資訊。
- 段落最後必須逐字使用：{conclusion}"""


def report_opening_instruction(test_item: str) -> str:
    """統一報告開場：先交代測項目的/範圍，再描述本次實際作法與結果。"""
    item = clean_text(test_item) or "(未提供)"
    return f"""【Report 開場規則】
- 第一句必須先說明此測試項目的評估目的或範圍，再進入本次實際測試方式與結果。
- 請依測試項目名稱「{item}」與測項類別，自然改寫成報告式開場，不要只重複測項名稱。
- 英文第一句優先使用類似：This test item evaluated ... for security vulnerabilities / security controls.
- 繁體中文第一句優先使用類似：本測試項目針對……進行安全弱點評估／安全性評估。
- 第一個句子只是在交代測項目的或評估範圍，不代表已使用特定工具或完成特定攻擊。
- 實際測試方式、工具、操作、觀察或確認方式，只能來自「目前情況」；若目前情況未提供，就不要自行補充。
- 不得因測項名稱自行宣稱已執行特定工具、payload、掃描或攻擊手法。
- Not Applicable 也必須先寫測項目的/範圍，再說明不適用或未執行原因。
- None / No Risk 先寫測項目的/範圍，再寫實際確認方式與未發現風險的結果。
- Low / Medium / High / Critical 先寫測項目的/範圍，再寫實際弱點、影響、限制條件與改善方向。
"""


REFERENCE_PATTERN = re.compile(r"(?:參考|参考|refer\s+to)\s*([A-Za-z]+\d+)", re.IGNORECASE)


def extract_reference_ids(text: str) -> list[str]:
    """從目前情況抓出如「參考P06」「參考 P06」「refer to P06」的測項編號。"""
    seen = set()
    refs = []
    for match in REFERENCE_PATTERN.finditer(clean_text(text)):
        ref = clean_text(match.group(1)).upper()
        if ref and ref not in seen:
            seen.add(ref)
            refs.append(ref)
    return refs


def build_reference_context(current_status: str, row_context_by_id: dict[str, dict]) -> str:
    """將被引用測項的既有資訊整理成可提供給 AI 的補充上下文。"""
    blocks = []
    for ref_id in extract_reference_ids(current_status):
        ref = row_context_by_id.get(ref_id)
        if not ref:
            continue
        block = (
            f"參考測項：{ref_id}\n"
            f"測試項目：{ref.get('test_item', '')}\n"
            f"判定結果：{ref.get('result', '')}\n"
            f"目前情況：\n{ref.get('current_status', '')}"
        )
        blocks.append(block)
    return "\n\n".join(blocks)


def build_prompt(
    *,
    item_no: str,
    test_item: str,
    result: str,
    current_status: str,
    remediation: str = "",
    language: str = "English",
    reference_context: str = "",
) -> str:
    # language 參數保留供既有 .env / 呼叫介面相容；實際固定輸出英文與繁體中文兩個欄位。
    remediation_text = clean_text(remediation) or "(未提供)"
    reference_text = clean_text(reference_context) or "(無)"
    structure_en = report_structure_instruction(result, "English")
    structure_zh = report_structure_instruction(result, "繁體中文")
    conclusion_en = fixed_conclusion(result, "English")
    conclusion_zh = fixed_conclusion(result, "繁體中文")
    item_category = classify_test_item(test_item)
    item_template = test_item_template_instruction(test_item)
    opening_rules = report_opening_instruction(test_item)
    kind = normalize_result_kind(result)

    if kind in {"low", "medium", "high", "critical"}:
        cvss_rule = """請另外判定 CVSS v3.1 Base Metrics：AV、AC、PR、UI、S、C、I、A。
- 主要依漏洞本身、「目前情況」以及下方明確列出的「參考測項內容」判定。
- 不要因環境特定緩解措施（例如額外防火牆、ACL、白名單）降低 Base Metrics。
- 請以 best-effort 完成八個 Base Metrics；不要只因某個 metric 未逐字寫明就直接回傳 null。
- 可依漏洞型態、攻擊路徑、是否需要鄰近/本地存取、是否需要使用者互動、以及實際影響做合理判斷。
- 不得為了讓 CVSS severity 配合人工的 Low/Medium/High/Critical 而刻意調整 metrics；人工判定與 CVSS Base severity 可以不同。
- 只有輸入內容完全不足以辨識漏洞或攻擊情境時，cvss 才可輸出 null。
- 不要自行計算或輸出 score；Python 會依 FIRST CVSS v3.1 公式計算。"""
    else:
        cvss_rule = "此判定不需要 AI 產生 Base Metrics，cvss 請輸出 null。"

    return f"""請根據以下檢測資料產生英文 Report、繁體中文 Report，以及需要時的 CVSS v3.1 Base Metrics。

測項編號：{clean_text(item_no)}
測試項目：{clean_text(test_item)}
判定結果：{clean_text(result)}

目前情況：
{clean_text(current_status)}

修補建議：
{remediation_text}

參考測項內容：
{reference_text}

測項模板代碼：{item_category}
{item_template}

【Report 統一開場規則】
{opening_rules}

【英文 Report 規則】
{structure_en}

【繁體中文 Report 規則】
{structure_zh}

【CVSS 規則】
{cvss_rule}

重要限制：
- 人工判定結果必須維持為 {clean_text(result)}，不得自行變更。
- 英文與繁體中文 Report 的第一句都必須先交代本測試項目的評估目的或範圍。
- 第一個句子之後，再依「目前情況」描述本次實際測試方式、確認方式、觀察結果或風險；不得把測項名稱本身當成已執行的證據。
- 英文固定結論句為：{conclusion_en}
- 繁體中文固定結論句為：{conclusion_zh}
- report_en 只放英文，不要加 English: 標籤。
- report_zh 只放繁體中文，不要加「繁體中文：」標籤。
- 兩份 Report 必須描述相同的測試事實、風險、降低風險因素及修補方向。
- 不得自行新增目前情況或「參考測項內容」中不存在的工具、弱點、產品能力、防護措施、限制條件或降低風險因素。
- 「參考測項內容」只能作為目前情況明確引用時的補充依據，不得改變本列人工判定結果。
- 若修補建議未提供，不要自行創造具體修補措施。
- 英文風險名稱一律小寫：not applicable、no risk、low risk、medium risk、high risk、critical risk。
- 僅輸出 JSON object，格式如下：
{{
  "report_en": "English report paragraph",
  "report_zh": "繁體中文報告段落",
  "cvss": {{"AV":"N","AC":"L","PR":"N","UI":"N","S":"U","C":"L","I":"N","A":"N"}}
}}
not applicable 與 no risk 請將 "cvss" 設為 null；low risk 以上原則上輸出完整八個 Base Metrics，只有資訊完全不足時才可設為 null。
"""



def _resolve_sheet_xml(zf: zipfile.ZipFile, sheet_name: str) -> str:
    workbook_root = ET.fromstring(zf.read("xl/workbook.xml"))
    ns_main = {"m": MAIN_NS}
    relationship_id = None

    for sheet in workbook_root.findall("m:sheets/m:sheet", ns_main):
        if sheet.attrib.get("name") == sheet_name:
            relationship_id = sheet.attrib.get(f"{{{REL_NS}}}id")
            break

    if not relationship_id:
        raise ValueError(f"找不到工作表：{sheet_name}")

    rels_root = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    target = None
    for rel in rels_root.findall(f"{{{PKG_REL_NS}}}Relationship"):
        if rel.attrib.get("Id") == relationship_id:
            target = rel.attrib.get("Target")
            break

    if not target:
        raise ValueError(f"無法解析工作表 {sheet_name} 的 XML 路徑")

    target_path = PurePosixPath(target)
    if target.startswith("/"):
        return str(target_path).lstrip("/")
    return str(PurePosixPath("xl") / target_path)


def extract_x14_list_validations(input_path: Path | str, sheet_name: str) -> list[dict]:
    """唯讀解析來源 XLSX 中的 x14 清單型 DataValidation。

    不修改 XML；只把 openpyxl 讀取時會遺失的規則先記錄下來，
    之後以標準 DataValidation 重建。
    """
    input_path = Path(input_path)
    validations: list[dict] = []

    with zipfile.ZipFile(input_path, "r") as zf:
        sheet_xml_path = _resolve_sheet_xml(zf, sheet_name)
        root = ET.fromstring(zf.read(sheet_xml_path))
        ns = {"x14": X14_NS, "xm": XM_NS}

        for dv in root.findall(".//x14:dataValidation", ns):
            dv_type = dv.attrib.get("type", "")
            formula_node = dv.find("x14:formula1/xm:f", ns)
            sqref_node = dv.find("xm:sqref", ns)

            if dv_type != "list":
                raise ValueError(
                    f"發現目前程式未支援轉換的 x14 DataValidation 類型：{dv_type or '(未指定)'}"
                )
            if formula_node is None or not clean_text(formula_node.text):
                raise ValueError("x14 清單型 DataValidation 缺少 formula1")
            if sqref_node is None or not clean_text(sqref_node.text):
                raise ValueError("x14 清單型 DataValidation 缺少 sqref")

            validations.append(
                {
                    "type": "list",
                    "formula": clean_text(formula_node.text),
                    "sqref": clean_text(sqref_node.text),
                    "allow_blank": dv.attrib.get("allowBlank", "0") == "1",
                    "show_error": dv.attrib.get("showErrorMessage", "0") == "1",
                    "show_input": dv.attrib.get("showInputMessage", "0") == "1",
                    "error_style": dv.attrib.get("errorStyle", "stop"),
                }
            )

    return validations


def _unique_defined_name(wb, base_name: str) -> str:
    if base_name not in wb.defined_names:
        return base_name
    i = 2
    while f"{base_name}_{i}" in wb.defined_names:
        i += 1
    return f"{base_name}_{i}"


def restore_x14_as_standard_validations(wb, ws, validations: list[dict]) -> int:
    """把先前讀出的 x14 清單驗證重建成標準 DataValidation。"""
    restored = 0

    for idx, spec in enumerate(validations, start=1):
        formula = spec["formula"]

        # 跨工作表的清單來源改用標準 Defined Name，Excel/openpyxl 相容性較好。
        if "!" in formula and not formula.startswith('"'):
            name = _unique_defined_name(wb, f"_AI_REPORT_DV_{idx}")
            wb.defined_names.add(DefinedName(name, attr_text=formula))
            formula1 = f"={name}"
        else:
            formula1 = formula if formula.startswith("=") or formula.startswith('"') else f"={formula}"

        dv = DataValidation(
            type="list",
            formula1=formula1,
            allow_blank=spec["allow_blank"],
            showErrorMessage=spec["show_error"],
            showInputMessage=spec["show_input"],
            errorStyle=spec["error_style"],
        )
        ws.add_data_validation(dv)

        # sqref 可能含多個區段，例如 D2:D25 F2:F25。
        for ref in spec["sqref"].split():
            dv.add(ref)

        restored += 1

    return restored


def find_header_map(ws) -> tuple[int, dict[str, int]]:
    for row_idx in range(1, min(20, ws.max_row) + 1):
        values = [clean_text(ws.cell(row_idx, col).value) for col in range(1, ws.max_column + 1)]
        if all(h in values for h in REQUIRED_HEADERS):
            headers = {value: idx + 1 for idx, value in enumerate(values) if value}
            return row_idx, headers
    raise ValueError("找不到必要欄位：" + ", ".join(REQUIRED_HEADERS))



RECOMMENDATION_SYSTEM_PROMPT = """你是一位資安檢測報告的改善建議撰寫人員。
你會收到同一個檢測大分類中，被人工判定為 low risk 或以上的測項。

請遵守：
1. recommendation_en 只輸出英文；recommendation_zh 只輸出繁體中文，兩者語意一致。
2. 建議必須是「大分類層級」的整體改善方向，不要逐項重複 Report。
3. 可以提出與已發現風險直接相關的資安最佳實務，但不得把未提供的產品功能、弱點、攻擊結果或既有控制措施寫成事實。
4. 使用建議語氣，例如 should consider, is recommended to, 建議、可評估；不要聲稱某個特定修補方式一定適用。
5. 若分類為 Static，應優先從 SSDLC、原始碼安全檢測、Code Review、SAST/SCA、SBOM/第三方元件弱點管理、修補追蹤與版本發布安全關卡等方向提出建議，但只選與觸發測項相關的內容。
6. 若分類為 Dynamic，應優先從動態安全測試、Runtime 行為、版本發布前回歸測試與異常行為驗證等方向提出建議。
7. 若分類為 Fuzzing，應優先從通訊協定、Parser、API、邊界輸入與異常資料的持續 Fuzz Testing 與回歸測試等方向提出建議。
8. 若分類為 PT，應依實際觸發測項聚焦攻擊面、通訊安全、身分驗證、授權、敏感資料保護、更新或實體安全等相關控制，不要把無關領域全部列入。
9. 每種語言輸出一個精簡正式段落，不要標題、Markdown、條列或前言。
10. 僅輸出 JSON object。

JSON 格式：
{
  "recommendation_en": "English recommendation paragraph",
  "recommendation_zh": "繁體中文建議段落"
}
"""

CATEGORY_RECOMMENDATION_HEADERS = ["分類", "最高風險", "觸發測項", "Recommendation", "Recommendation_ch"]
RISK_RANK = {"not_applicable": 0, "none": 1, "low": 2, "medium": 3, "high": 4, "critical": 5}


def collect_category_summaries(ws, header_row: int, headers: dict[str, int]) -> list[dict]:
    """依「分類」彙整測項；只有 low risk 以上會列入 triggered_items。"""
    category_col = headers.get("分類")
    if not category_col:
        raise ValueError("找不到欄位：分類")

    remediation_col = headers.get("修補建議")
    report_col = headers.get("Report")
    report_ch_col = headers.get("Report_ch")
    categories: dict[str, dict] = {}

    for row_idx in range(header_row + 1, ws.max_row + 1):
        category = clean_text(ws.cell(row_idx, category_col).value)
        item_no = clean_text(ws.cell(row_idx, headers["編號"]).value)
        test_item = clean_text(ws.cell(row_idx, headers["測項"]).value)
        result = clean_text(ws.cell(row_idx, headers["判定結果"]).value)
        current_status = clean_text(ws.cell(row_idx, headers["目前情況"]).value)
        if not category or not any([item_no, test_item, result, current_status]):
            continue

        summary = categories.setdefault(category, {
            "category": category,
            "highest_risk": "-",
            "triggered_items": [],
            "_highest_rank": -1,
            "_has_pending": False,
        })

        kind = normalize_result_kind(result)
        if result.casefold() in SKIP_RESULTS or not result:
            summary["_has_pending"] = True
            continue
        if kind not in RISK_RANK:
            continue

        rank = RISK_RANK[kind]
        if rank > summary["_highest_rank"]:
            summary["_highest_rank"] = rank
            summary["highest_risk"] = RESULT_LABEL_EN[kind]

        if kind in {"low", "medium", "high", "critical"}:
            summary["triggered_items"].append({
                "item_no": item_no,
                "test_item": test_item,
                "result": result,
                "current_status": current_status,
                "remediation": clean_text(ws.cell(row_idx, remediation_col).value) if remediation_col else "",
                "report_en": clean_text(ws.cell(row_idx, report_col).value) if report_col else "",
                "report_zh": clean_text(ws.cell(row_idx, report_ch_col).value) if report_ch_col else "",
            })

    output = []
    for summary in categories.values():
        # 只有 N/A 且仍有 TBD/Testing 時，最高風險顯示 -，避免看起來像已完成整類評估。
        if summary["_highest_rank"] <= 0 and summary["_has_pending"]:
            summary["highest_risk"] = "-"
        summary.pop("_highest_rank", None)
        summary.pop("_has_pending", None)
        output.append(summary)
    return output


def build_category_recommendation_prompt(category: str, triggered_items: list[dict]) -> str:
    lines = []
    for item in triggered_items:
        lines.append(
            f"- {item['item_no']} | {item['test_item']} | 判定結果：{item['result']}\n"
            f"  目前情況：{item['current_status'] or '(未提供)'}\n"
            f"  原有修補建議：{item['remediation'] or '(未提供)'}"
        )
    findings = "\n".join(lines) if lines else "(無 low risk 以上測項)"
    return f"""請針對以下資安檢測大分類，提出給廠商的整體改善建議。

分類：{category}

觸發測項（僅列 low risk 以上）：
{findings}

要求：
- 不要逐字重述每一筆 finding，而要將同類風險整理成後續開發、維運或安全驗證可採取的整體改善方向。
- 建議必須對應上述觸發測項；可引用相關資安最佳實務，但不要臆測產品目前已具備或缺少未提及的功能。
- recommendation_en 與 recommendation_zh 必須語意一致。
- 僅輸出 JSON object。
"""


def parse_recommendation_payload(text: str) -> dict:
    raw = clean_text(text)
    if raw.startswith("```"):
        raw = raw.removeprefix("```json").removeprefix("```").strip()
        if raw.endswith("```"):
            raw = raw[:-3].strip()
    first = raw.find("{")
    last = raw.rfind("}")
    if first < 0 or last < first:
        raise ValueError("建議 AI 回傳內容不是有效 JSON object")
    try:
        payload = json.loads(raw[first:last + 1])
    except json.JSONDecodeError as exc:
        raise ValueError(f"建議 AI JSON 解析失敗：{exc}") from exc
    recommendation_en = clean_text(payload.get("recommendation_en"))
    recommendation_zh = clean_text(payload.get("recommendation_zh"))
    if not recommendation_en or not recommendation_zh:
        raise ValueError("建議 AI JSON 缺少 recommendation_en 或 recommendation_zh")
    return {"recommendation_en": recommendation_en, "recommendation_zh": recommendation_zh}


def make_recommendation_generator() -> Callable[..., str]:
    provider = os.getenv("AI_PROVIDER", "openai").strip().lower()

    if provider == "openai":
        try:
            from openai import OpenAI
        except ImportError as exc:
            raise RuntimeError("尚未安裝 openai，請執行：pip install openai") from exc
        api_key = os.getenv("OPENAI_API_KEY", "").strip()
        model = os.getenv("OPENAI_MODEL", "").strip()
        reasoning_effort = os.getenv("OPENAI_REASONING", "").strip()
        if not api_key:
            raise RuntimeError(".env 尚未設定 OPENAI_API_KEY")
        if not model:
            raise RuntimeError(".env 尚未設定 OPENAI_MODEL")
        client = OpenAI(api_key=api_key)

        def generate_openai_recommendation(**kwargs) -> str:
            request = {
                "model": model,
                "instructions": RECOMMENDATION_SYSTEM_PROMPT,
                "input": build_category_recommendation_prompt(**kwargs),
                "store": False,
            }
            if reasoning_effort:
                request["reasoning"] = {"effort": reasoning_effort}
            response = client.responses.create(**request)
            text = clean_text(response.output_text)
            if not text:
                raise RuntimeError("OpenAI API 未回傳分類建議")
            return text
        return generate_openai_recommendation

    if provider == "gemini":
        try:
            from google import genai
            from google.genai import types
        except ImportError as exc:
            raise RuntimeError("尚未安裝 google-genai，請執行：pip install google-genai") from exc
        api_key = os.getenv("GEMINI_API_KEY", "").strip()
        model = os.getenv("GEMINI_MODEL", "").strip()
        if not api_key:
            raise RuntimeError(".env 尚未設定 GEMINI_API_KEY")
        if not model:
            raise RuntimeError(".env 尚未設定 GEMINI_MODEL")
        client = genai.Client(api_key=api_key)

        def generate_gemini_recommendation(**kwargs) -> str:
            response = call_gemini_generate_content(
                client.models.generate_content,
                model=model,
                contents=build_category_recommendation_prompt(**kwargs),
                config=types.GenerateContentConfig(
                    system_instruction=RECOMMENDATION_SYSTEM_PROMPT,
                    temperature=0.2,
                ),
            )
            text = clean_text(response.text)
            if not text:
                raise RuntimeError("Gemini API 未回傳分類建議")
            return text
        return generate_gemini_recommendation

    if provider == "ollama":
        backend = get_ollama_backend()

        def generate_ollama_recommendation(**kwargs) -> str:
            return backend.chat(
                system_prompt=RECOMMENDATION_SYSTEM_PROMPT,
                user_prompt=build_category_recommendation_prompt(**kwargs),
                schema=RECOMMENDATION_RESPONSE_SCHEMA,
                temperature=0.2,
                label=f"Recommendation:{clean_text(kwargs.get('category'))}",
            )

        return generate_ollama_recommendation

    raise RuntimeError("AI_PROVIDER 僅支援 openai、gemini 或 ollama")


def write_security_recommendation_sheet(wb, summaries: list[dict], recommendation_generate_func: Callable[..., str] | None) -> dict:
    """建立/重建 Security Recommendation 工作表。"""
    sheet_name = "Security Recommendation"
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    insert_index = 1 if "IOT Device" in wb.sheetnames else len(wb.sheetnames)
    ws = wb.create_sheet(sheet_name, insert_index)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    body_alignment = Alignment(vertical="top", wrap_text=True)

    for col_idx, header in enumerate(CATEGORY_RECOMMENDATION_HEADERS, start=1):
        cell = ws.cell(1, col_idx, header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    generated = 0
    failed = 0
    for row_idx, summary in enumerate(summaries, start=2):
        ws.cell(row_idx, 1).value = summary["category"]
        ws.cell(row_idx, 2).value = summary["highest_risk"]
        triggered = summary["triggered_items"]
        ws.cell(row_idx, 3).value = ", ".join(item["item_no"] for item in triggered) if triggered else "-"

        if triggered and recommendation_generate_func is not None:
            try:
                raw = recommendation_generate_func(category=summary["category"], triggered_items=triggered)
                payload = parse_recommendation_payload(raw)
                ws.cell(row_idx, 4).value = payload["recommendation_en"]
                ws.cell(row_idx, 5).value = payload["recommendation_zh"]
                generated += 1
            except Exception as exc:
                failed += 1
                print(f"  [RECOMMENDATION ERROR] {summary['category']}: {exc}", file=sys.stderr, flush=True)

        for col_idx in range(1, 6):
            ws.cell(row_idx, col_idx).alignment = body_alignment
        for col_idx in (1, 2, 3):
            ws.cell(row_idx, col_idx).alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)

        # Recommendation 長度每案不同，依中英文內容自動增加列高。
        auto_adjust_row_height(
            ws, row_idx, (4, 5),
            min_height=30.0, max_height=405.0, line_height=15.0, padding=8.0,
        )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:E{max(ws.max_row, 1)}"
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["D"].width = 70
    ws.column_dimensions["E"].width = 70
    ws.row_dimensions[1].height = 28
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    return {"recommendation_generated": generated, "recommendation_failed": failed}

def make_generator() -> Callable[..., str]:
    provider = os.getenv("AI_PROVIDER", "openai").strip().lower()

    if provider == "openai":
        try:
            from openai import OpenAI
        except ImportError as exc:
            raise RuntimeError("尚未安裝 openai，請執行：pip install openai") from exc

        api_key = os.getenv("OPENAI_API_KEY", "").strip()
        model = os.getenv("OPENAI_MODEL", "").strip()
        reasoning_effort = os.getenv("OPENAI_REASONING", "").strip()

        if not api_key:
            raise RuntimeError(".env 尚未設定 OPENAI_API_KEY")
        if not model:
            raise RuntimeError(".env 尚未設定 OPENAI_MODEL")

        client = OpenAI(api_key=api_key)

        def generate_openai(**kwargs) -> str:
            request = {
                "model": model,
                "instructions": SYSTEM_PROMPT,
                "input": build_prompt(**kwargs),
                "store": False,
            }
            if reasoning_effort:
                request["reasoning"] = {"effort": reasoning_effort}

            response = client.responses.create(**request)
            text = clean_text(response.output_text)
            if not text:
                raise RuntimeError("OpenAI API 未回傳文字內容")
            return text

        return generate_openai

    if provider == "gemini":
        try:
            from google import genai
            from google.genai import types
        except ImportError as exc:
            raise RuntimeError("尚未安裝 google-genai，請執行：pip install google-genai") from exc

        api_key = os.getenv("GEMINI_API_KEY", "").strip()
        model = os.getenv("GEMINI_MODEL", "").strip()

        if not api_key:
            raise RuntimeError(".env 尚未設定 GEMINI_API_KEY")
        if not model:
            raise RuntimeError(".env 尚未設定 GEMINI_MODEL")

        client = genai.Client(api_key=api_key)

        def generate_gemini(**kwargs) -> str:
            response = call_gemini_generate_content(
                client.models.generate_content,
                model=model,
                contents=build_prompt(**kwargs),
                config=types.GenerateContentConfig(
                    system_instruction=SYSTEM_PROMPT,
                    temperature=0.2,
                ),
            )
            text = clean_text(response.text)
            if not text:
                raise RuntimeError("Gemini API 未回傳文字內容")
            return text

        return generate_gemini

    if provider == "ollama":
        backend = get_ollama_backend()

        def generate_ollama(**kwargs) -> str:
            return backend.chat(
                system_prompt=SYSTEM_PROMPT,
                user_prompt=build_prompt(**kwargs),
                schema=REPORT_RESPONSE_SCHEMA,
                temperature=0.2,
                label=clean_text(kwargs.get("item_no")),
            )

        return generate_ollama

    raise RuntimeError("AI_PROVIDER 僅支援 openai、gemini 或 ollama")


def make_cvss_retry_generator() -> Callable[..., str]:
    """建立僅用於 CVSS 第二次判讀的 API 呼叫器。"""
    provider = os.getenv("AI_PROVIDER", "openai").strip().lower()

    if provider == "openai":
        try:
            from openai import OpenAI
        except ImportError as exc:
            raise RuntimeError("尚未安裝 openai，請執行：pip install openai") from exc
        api_key = os.getenv("OPENAI_API_KEY", "").strip()
        model = os.getenv("OPENAI_MODEL", "").strip()
        reasoning_effort = os.getenv("OPENAI_REASONING", "").strip()
        if not api_key:
            raise RuntimeError(".env 尚未設定 OPENAI_API_KEY")
        if not model:
            raise RuntimeError(".env 尚未設定 OPENAI_MODEL")
        client = OpenAI(api_key=api_key)

        def generate_openai_cvss_retry(**kwargs) -> str:
            request = {
                "model": model,
                "instructions": CVSS_RETRY_SYSTEM_PROMPT,
                "input": build_cvss_retry_prompt(**kwargs),
                "store": False,
            }
            if reasoning_effort:
                request["reasoning"] = {"effort": reasoning_effort}
            response = client.responses.create(**request)
            text = clean_text(response.output_text)
            if not text:
                raise RuntimeError("OpenAI API 未回傳 CVSS 重試內容")
            return text

        return generate_openai_cvss_retry

    if provider == "gemini":
        try:
            from google import genai
            from google.genai import types
        except ImportError as exc:
            raise RuntimeError("尚未安裝 google-genai，請執行：pip install google-genai") from exc
        api_key = os.getenv("GEMINI_API_KEY", "").strip()
        model = os.getenv("GEMINI_MODEL", "").strip()
        if not api_key:
            raise RuntimeError(".env 尚未設定 GEMINI_API_KEY")
        if not model:
            raise RuntimeError(".env 尚未設定 GEMINI_MODEL")
        client = genai.Client(api_key=api_key)

        def generate_gemini_cvss_retry(**kwargs) -> str:
            response = call_gemini_generate_content(
                client.models.generate_content,
                model=model,
                contents=build_cvss_retry_prompt(**kwargs),
                config=types.GenerateContentConfig(
                    system_instruction=CVSS_RETRY_SYSTEM_PROMPT,
                    temperature=0.1,
                ),
            )
            text = clean_text(response.text)
            if not text:
                raise RuntimeError("Gemini API 未回傳 CVSS 重試內容")
            return text

        return generate_gemini_cvss_retry

    if provider == "ollama":
        backend = get_ollama_backend()

        def generate_ollama_cvss_retry(**kwargs) -> str:
            return backend.chat(
                system_prompt=CVSS_RETRY_SYSTEM_PROMPT,
                user_prompt=build_cvss_retry_prompt(**kwargs),
                schema=CVSS_RETRY_RESPONSE_SCHEMA,
                temperature=0.1,
                label=f"CVSS:{clean_text(kwargs.get('item_no'))}",
            )

        return generate_ollama_cvss_retry

    raise RuntimeError("AI_PROVIDER 僅支援 openai、gemini 或 ollama")


def process_workbook(
    *,
    input_path: Path | str,
    output_path: Path | str,
    sheet_name: str,
    overwrite_report: bool,
    language: str,
    generate_func: Callable[..., str] | None = None,
    recommendation_generate_func: Callable[..., str] | None = None,
    cvss_retry_func: Callable[..., str] | None = None,
    normalize_only: bool = False,
) -> dict:
    input_path = Path(input_path)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    # 先從原始 xlsx 記住 x14 DataValidation；openpyxl 載入後會移除它。
    x14_validations = extract_x14_list_validations(input_path, sheet_name)

    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="Data Validation extension is not supported and will be removed",
            category=UserWarning,
        )
        wb = load_workbook(input_path, data_only=False)

    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"找不到工作表：{sheet_name}")
        ws = wb[sheet_name]

        restored_validations = restore_x14_as_standard_validations(wb, ws, x14_validations)
        header_row, headers = find_header_map(ws)

        remediation_col = headers.get("修補建議")
        report_col = headers["Report"]
        report_ch_col = headers["Report_ch"]
        cvss_col = headers["cvss"]
        score_col = headers["score"]
        if not normalize_only:
            generate_func = generate_func or make_generator()
            recommendation_generate_func = recommendation_generate_func or make_recommendation_generator()

        generated = 0
        skipped = 0
        failed = 0
        cvss_failed = 0
        cvss_na = 0

        # 建立測項索引，讓「參考P06」這類文字能取得被引用列的上下文。
        row_context_by_id = {}
        for context_row_idx in range(header_row + 1, ws.max_row + 1):
            context_item_no = clean_text(ws.cell(context_row_idx, headers["編號"]).value).upper()
            if not context_item_no:
                continue
            row_context_by_id[context_item_no] = {
                "test_item": clean_text(ws.cell(context_row_idx, headers["測項"]).value),
                "result": clean_text(ws.cell(context_row_idx, headers["判定結果"]).value),
                "current_status": clean_text(ws.cell(context_row_idx, headers["目前情況"]).value),
            }

        for row_idx in range(header_row + 1, ws.max_row + 1):
            item_no = clean_text(ws.cell(row_idx, headers["編號"]).value)
            test_item = clean_text(ws.cell(row_idx, headers["測項"]).value)
            result = clean_text(ws.cell(row_idx, headers["判定結果"]).value)
            current_status = clean_text(ws.cell(row_idx, headers["目前情況"]).value)
            remediation = clean_text(ws.cell(row_idx, remediation_col).value) if remediation_col else ""
            existing_report = clean_text(ws.cell(row_idx, report_col).value)
            existing_report_ch = clean_text(ws.cell(row_idx, report_ch_col).value)
            existing_cvss = clean_text(ws.cell(row_idx, cvss_col).value)
            existing_score = clean_text(ws.cell(row_idx, score_col).value)

            if not any([
                item_no, test_item, result, current_status, remediation,
                existing_report, existing_report_ch, existing_cvss, existing_score,
            ]):
                continue

            if normalize_only:
                continue

            if not result or not current_status or result.casefold() in SKIP_RESULTS:
                skipped += 1
                continue

            if not overwrite_report and all([
                existing_report, existing_report_ch, existing_cvss, existing_score,
            ]):
                skipped += 1
                continue

            try:
                print(f"[{row_idx}] {item_no} {test_item} ({result}) ...", flush=True)
                reference_context = build_reference_context(current_status, row_context_by_id)
                raw_ai = generate_func(
                    item_no=item_no,
                    test_item=test_item,
                    result=result,
                    current_status=current_status,
                    remediation=remediation,
                    language=language,
                    reference_context=reference_context,
                )
                payload = parse_ai_payload(raw_ai)

                # Report 與 CVSS 分開處理：CVSS 判讀失敗不應吃掉已成功產生的雙語報告。
                ws.cell(row_idx, report_col).value = payload["report_en"]
                ws.cell(row_idx, report_ch_col).value = payload["report_zh"]
                generated += 1

                try:
                    vector, score = cvss_for_result(result, payload.get("cvss"))
                except Exception as first_cvss_exc:
                    kind = normalize_result_kind(result)
                    if kind in {"low", "medium", "high", "critical"}:
                        try:
                            if cvss_retry_func is None:
                                cvss_retry_func = make_cvss_retry_generator()
                            print(
                                f"  [CVSS RETRY] {item_no}: 第一輪 CVSS 不完整，重新判讀一次",
                                file=sys.stderr, flush=True,
                            )
                            retry_raw = cvss_retry_func(
                                item_no=item_no,
                                test_item=test_item,
                                result=result,
                                current_status=current_status,
                                remediation=remediation,
                                reference_context=reference_context,
                            )
                            retry_metrics = parse_cvss_retry_payload(retry_raw)
                            if retry_metrics is None:
                                cvss_na += 1
                                vector, score = "N/A", "-"
                                print(
                                    f"  [CVSS N/A] {item_no}: 目前資訊不足以辨識可供 CVSS v3.1 Base Metrics 評估的具體漏洞或攻擊情境",
                                    file=sys.stderr, flush=True,
                                )
                            else:
                                vector, score = cvss_for_result(result, retry_metrics)
                        except Exception as retry_exc:
                            cvss_failed += 1
                            print(
                                f"  [CVSS ERROR] {item_no}: 第一輪：{first_cvss_exc}；重試：{retry_exc}",
                                file=sys.stderr, flush=True,
                            )
                            vector = score = None
                    else:
                        cvss_failed += 1
                        print(f"  [CVSS ERROR] {item_no}: {first_cvss_exc}", file=sys.stderr, flush=True)
                        vector = score = None

                if vector is not None:
                    cvss_cell = ws.cell(row_idx, cvss_col)
                    score_cell = ws.cell(row_idx, score_col)
                    cvss_cell.value = vector
                    score_cell.value = score

                    # 數值型 score 固定顯示一位小數，例如 None/no risk 顯示為 0.0。
                    if isinstance(score, (int, float)):
                        score_cell.number_format = "0.0"
                        warning = score_result_mismatch_message(result, float(score))
                        mismatch = bool(warning)
                        set_cvss_cells_mismatch_style(ws, row_idx, cvss_col, score_col, mismatch)
                        if warning:
                            print(f"  [CVSS WARNING] {item_no}: {warning}", file=sys.stderr, flush=True)
            except Exception as exc:
                failed += 1
                print(f"  [ERROR] {item_no}: {exc}", file=sys.stderr, flush=True)

        # IOT Device 的文字欄位依實際內容自動增加列高。
        # 包含未呼叫 AI 的列，避免「目前情況／修補建議」本身過長而被截斷。
        if not normalize_only:
            text_columns = [
                headers.get("目前情況"),
                headers.get("修補建議"),
                headers.get("Report"),
                headers.get("Report_ch"),
            ]
            text_columns = [col for col in text_columns if col]
            for auto_row_idx in range(header_row + 1, ws.max_row + 1):
                if any(clean_text(ws.cell(auto_row_idx, col).value) for col in text_columns):
                    auto_adjust_row_height(
                        ws, auto_row_idx, text_columns,
                        min_height=20.0, max_height=405.0, line_height=18.0, padding=8.0,
                    )

        # 自動列高完成後，再統一四個輸出欄位格式，避免列高函式把 Report 對齊改回靠上。
        # cvss / score 只調整字級與對齊，保留既有字色，因此 mismatch 紅字不會被覆蓋。
        if not normalize_only:
            for format_row_idx in range(header_row + 1, ws.max_row + 1):
                format_output_cells(
                    ws, format_row_idx,
                    report_col, report_ch_col, cvss_col, score_col,
                )

        recommendation_stats = {"recommendation_generated": 0, "recommendation_failed": 0}
        if not normalize_only:
            summaries = collect_category_summaries(ws, header_row, headers)
            recommendation_stats = write_security_recommendation_sheet(
                wb, summaries, recommendation_generate_func
            )

        wb.save(output_path)
    finally:
        wb.close()

    verify_wb = load_workbook(output_path, read_only=True, data_only=False)
    verify_wb.close()

    return {
        "generated": generated,
        "skipped": skipped,
        "failed": failed,
        "cvss_failed": cvss_failed,
        "cvss_na": cvss_na,
        "restored_validations": restored_validations,
        "normalize_only": normalize_only,
        **recommendation_stats,
        "output": str(output_path),
    }


def main() -> int:
    try:
        from dotenv import load_dotenv
    except ImportError:
        print("尚未安裝 python-dotenv，請執行：pip install python-dotenv", file=sys.stderr)
        return 1

    env_path = Path(__file__).resolve().with_name(".env")
    load_dotenv(env_path, override=True)

    provider = os.getenv("AI_PROVIDER", "openai").strip().lower()
    print(f"[CONFIG] .env: {env_path}")
    print(f"[AI] Provider: {provider}")
    if provider == "ollama":
        print(f"[AI] Model: {os.getenv('OLLAMA_MODEL', '').strip()}")
        print(f"[AI] Host: {os.getenv('OLLAMA_HOST', '').strip()}")
        print(f"[AI] Context: {os.getenv('OLLAMA_NUM_CTX', '8192').strip()}")
        print(f"[AI] Thinking: {os.getenv('OLLAMA_THINK', 'false').strip()}")
    elif provider == "gemini":
        print(f"[AI] Model: {os.getenv('GEMINI_MODEL', '').strip()}")
    elif provider == "openai":
        print(f"[AI] Model: {os.getenv('OPENAI_MODEL', '').strip()}")

    parser = argparse.ArgumentParser(description="AI 自動產生 Excel 資安檢測 Report")
    parser.add_argument("excel_path", help="輸入 Excel (.xlsx) 路徑")
    parser.add_argument("--output", help="輸出 Excel 路徑；未指定時自動加上 _report")
    parser.add_argument(
        "--normalize-only",
        action="store_true",
        help="只將 x14 下拉選單轉成標準 DataValidation，不呼叫 AI、不產生 Report",
    )
    args = parser.parse_args()

    input_path = Path(args.excel_path).expanduser().resolve()
    if not input_path.exists():
        print(f"找不到檔案：{input_path}", file=sys.stderr)
        return 1
    if input_path.suffix.lower() != ".xlsx":
        print("目前僅支援 .xlsx 檔案", file=sys.stderr)
        return 1

    output_path = (
        Path(args.output).expanduser().resolve()
        if args.output
        else input_path.with_name(
            f"{input_path.stem}_template{input_path.suffix}"
            if args.normalize_only
            else f"{input_path.stem}_report{input_path.suffix}"
        )
    )
    if output_path == input_path:
        print("輸出檔不可與原始檔相同，避免覆蓋原始檢測資料。", file=sys.stderr)
        return 1

    sheet_name = os.getenv("SHEET_NAME", "IOT Device").strip()
    language = os.getenv("REPORT_LANGUAGE", "English").strip()
    overwrite_report = str_to_bool(os.getenv("OVERWRITE_REPORT", "true"), True)

    try:
        stats = process_workbook(
            input_path=input_path,
            output_path=output_path,
            sheet_name=sheet_name,
            overwrite_report=overwrite_report,
            language=language,
            normalize_only=args.normalize_only,
        )
    except Exception as exc:
        print(f"執行失敗：{exc}", file=sys.stderr)
        return 1
    finally:
        close_ollama_backend()

    print("\n完成")
    if stats["normalize_only"]:
        print("模式：僅標準化 Excel，不呼叫 AI")
    print(f"產生 Report：{stats['generated']} 筆")
    print(f"略過：{stats['skipped']} 筆")
    print(f"失敗：{stats['failed']} 筆")
    if stats.get("cvss_na"):
        print(f"CVSS 不適用／資訊不足：{stats['cvss_na']} 筆")
    if stats.get("cvss_failed"):
        print(f"CVSS 需人工確認：{stats['cvss_failed']} 筆")
    if stats.get("recommendation_generated"):
        print(f"產生分類建議：{stats['recommendation_generated']} 類")
    if stats.get("recommendation_failed"):
        print(f"分類建議失敗：{stats['recommendation_failed']} 類")
    if stats["restored_validations"]:
        print(f"轉換 x14 下拉選單：{stats['restored_validations']} 組")
    print(f"輸出檔：{stats['output']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
